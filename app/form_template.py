"""產生對應 QWF-ME061 欄位的空白登打輔助表（非正式複印件）。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from .schema import FormSchema


def create_blank_form(schema: FormSchema, output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "登打輔助"

    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    title = f"{schema.form_name} ({schema.form_id} {schema.form_version})".strip()
    ws.merge_cells("A1:D1")
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    row = 3
    for field in schema.header_fields:
        ws.cell(row=row, column=1, value=field.label).font = Font(bold=True)
        ws.cell(row=row, column=2, value="")
        for c in range(1, 3):
            ws.cell(row=row, column=c).border = thin
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="分類").font = Font(bold=True)
    ws.cell(row=row, column=2, value="代碼").font = Font(bold=True)
    ws.cell(row=row, column=3, value="名稱").font = Font(bold=True)
    ws.cell(row=row, column=4, value="不良數").font = Font(bold=True)
    for c in range(1, 5):
        ws.cell(row=row, column=c).border = thin

    for item in schema.defect_items:
        row += 1
        ws.cell(row=row, column=1, value=item.category).border = thin
        ws.cell(row=row, column=2, value=item.key).border = thin
        ws.cell(row=row, column=3, value=item.label).border = thin
        ws.cell(row=row, column=4, value="").border = thin

    tip = row + 2
    ws.merge_cells(start_row=tip, start_column=1, end_row=tip, end_column=4)
    ws.cell(
        row=tip,
        column=1,
        value=(
            "拍照提示：整張入鏡、壓平、避免反光；底部「不良數」列與總計務必清晰。"
            "此檔為登打輔助清單，非正式紙本複印件。"
        ),
    )

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 12
    wb.save(output_path)
    return output_path
