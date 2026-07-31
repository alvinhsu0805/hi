"""把離線表頭辨識結果寫入 Excel。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .field_parsers import HeaderFields


HEADERS = [
    ("model_no", "型號"),
    ("lot_no", "批號(原文)"),
    ("lot_normalized", "批號(正規化)"),
    ("operator", "作業人員"),
    ("source_image", "來源影像"),
    ("warnings", "提醒"),
]


class HeaderExcelWriter:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)

    def ensure(self) -> None:
        if self.path.exists():
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "表頭辨識"
        fill = PatternFill("solid", fgColor="1F4E79")
        font = Font(color="FFFFFF", bold=True)
        for idx, (_, label) in enumerate(HEADERS, start=1):
            cell = ws.cell(1, idx, label)
            cell.fill = fill
            cell.font = font
        wb.save(self.path)

    def append(self, fields: HeaderFields, source_image: str = "") -> int:
        self.ensure()
        wb = load_workbook(self.path)
        ws = wb.active
        row = ws.max_row + 1
        values = {
            "model_no": fields.model_no,
            "lot_no": fields.lot_no,
            "lot_normalized": fields.lot_normalized,
            "operator": fields.operator,
            "source_image": source_image,
            "warnings": "；".join(fields.warnings),
        }
        for idx, (key, _) in enumerate(HEADERS, start=1):
            ws.cell(row, idx, values.get(key, ""))
        wb.save(self.path)
        return row
