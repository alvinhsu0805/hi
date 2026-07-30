"""將辨識結果寫入 Excel 不良統計表。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .schema import FormReadResult, FormSchema


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")


COLUMN_LABELS = {
    "date": "日期",
    "work_order": "製令",
    "product_no": "品號",
    "operator": "作業人員",
    "shift": "班別",
    "total": "不良合計",
    "source_image": "來源影像",
    "confidence": "信心分數",
    "reviewed": "已審核",
}


class ExcelWriter:
    def __init__(self, schema: FormSchema, workbook_path: str | Path):
        self.schema = schema
        self.path = Path(workbook_path)
        self.path.parent.mkdir(parents=True, exist_ok=True)

    def ensure_workbook(self) -> None:
        if self.path.exists():
            return
        wb = Workbook()
        ws = wb.active
        ws.title = self.schema.excel.sheet_name
        labels = self._header_labels()
        for key, col in self.schema.excel.columns.items():
            cell = ws[f"{col}{self.schema.excel.header_row}"]
            cell.value = labels.get(key, key)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        wb.save(self.path)

    def _header_labels(self) -> dict[str, str]:
        labels = dict(COLUMN_LABELS)
        for item in self.schema.defect_items:
            labels[item.key] = item.label
        return labels

    def append_result(self, result: FormReadResult, reviewed: bool = False) -> int:
        self.ensure_workbook()
        wb = load_workbook(self.path)
        ws = wb[self.schema.excel.sheet_name]
        row = ws.max_row + 1
        if ws.max_row == self.schema.excel.header_row and _row_empty(ws, row - 1):
            # header only
            pass

        values = {
            "date": result.date,
            "work_order": result.work_order,
            "product_no": result.product_no,
            "operator": result.operator,
            "shift": result.shift,
            "total": result.total_defects,
            "source_image": result.source_image,
            "confidence": round(result.overall_confidence, 3),
            "reviewed": "Y" if reviewed else "N",
        }
        for defect in result.defects:
            values[defect.key] = defect.count

        low_conf = result.overall_confidence < 0.7 or any(
            d.confidence < 0.6 for d in result.defects
        )
        for key, col in self.schema.excel.columns.items():
            cell = ws[f"{col}{row}"]
            cell.value = values.get(key, "")
            if low_conf and not reviewed:
                cell.fill = WARNING_FILL

        wb.save(self.path)
        return row


def _row_empty(ws, row: int) -> bool:
    for cell in ws[row]:
        if cell.value not in (None, ""):
            return False
    return True
